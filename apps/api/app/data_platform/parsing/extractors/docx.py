"""
Microsoft Word DOCX parser for WaterDIP.
"""

from __future__ import annotations

from io import BytesIO

from docx import Document

from ..base import BaseParser, ParseSource, ParserContext
from ..models import (
    DocumentSection,
    DocumentTable,
    DocumentType,
    WaterDIPDocument,
)
from .common import (
    build_source_provenance,
    deterministic_id,
    filename_title,
    normalise_text,
)


class DOCXParser(BaseParser[WaterDIPDocument]):
    """Parse Microsoft Word DOCX engineering documents."""

    name = "docx"
    parser_version = "1.0"

    supported_extensions = frozenset({".docx"})

    supported_media_types = frozenset(
        {
            (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        }
    )

    priority = 30

    def parse(
        self,
        source: ParseSource,
        context: ParserContext,
    ) -> WaterDIPDocument:
        content = source.read_bytes()
        document = Document(BytesIO(content))

        checksum = source.checksum()

        sections: list[DocumentSection] = []
        tables: list[DocumentTable] = []
        text_parts: list[str] = []

        current_title: str | None = None
        current_level = 1
        current_lines: list[str] = []
        section_index = 0

        def flush_section() -> None:
            nonlocal section_index
            nonlocal current_lines
            nonlocal current_title
            nonlocal current_level

            text = normalise_text(
                "\n".join(current_lines)
            )

            if not text and not current_title:
                current_lines = []
                return

            sections.append(
                DocumentSection(
                    id=deterministic_id(
                        "wd-section",
                        checksum,
                        section_index,
                        current_title or "",
                    ),
                    title=current_title,
                    level=current_level,
                    text=text,
                    order=section_index,
                )
            )

            section_index += 1
            current_lines = []

        for paragraph in document.paragraphs:
            text = normalise_text(paragraph.text)

            if not text:
                continue

            style_name = (
                paragraph.style.name
                if paragraph.style is not None
                else ""
            )

            if style_name.lower().startswith("heading"):
                flush_section()

                current_title = text

                try:
                    current_level = int(
                        style_name.split()[-1]
                    )
                except (
                    ValueError,
                    IndexError,
                ):
                    current_level = 1

                text_parts.append(text)

            else:
                current_lines.append(text)
                text_parts.append(text)

        flush_section()

        for table_index, table in enumerate(
            document.tables
        ):
            raw_rows = [
                [
                    normalise_text(cell.text)
                    for cell in row.cells
                ]
                for row in table.rows
            ]

            if not raw_rows:
                continue

            columns = raw_rows[0]
            rows = raw_rows[1:]

            canonical_table = DocumentTable(
                id=deterministic_id(
                    "wd-table",
                    checksum,
                    table_index,
                ),
                title=f"Table {table_index + 1}",
                columns=columns,
                rows=rows,
                metadata={
                    "table_index": table_index,
                    "row_count": len(rows),
                    "column_count": len(columns),
                },
            )

            tables.append(canonical_table)

            text_parts.append(
                f"Table {table_index + 1}"
            )
            text_parts.append(
                " | ".join(columns)
            )

            for row in rows:
                text_parts.append(
                    " | ".join(row)
                )

        title = None

        core_title = document.core_properties.title

        if core_title:
            title = core_title.strip() or None

        title = title or filename_title(source.filename)

        provenance = build_source_provenance(
            source,
            context,
            parser_name=self.name,
            parser_version=self.parser_version,
            checksum=checksum,
            metadata={
                "paragraph_count": len(document.paragraphs),
                "table_count": len(tables),
            },
        )

        return WaterDIPDocument.create(
            source=provenance,
            title=title,
            text="\n\n".join(text_parts),
            document_type=DocumentType.DOCUMENT,
            sections=sections,
            tables=tables,
            metadata={
                "paragraph_count": len(document.paragraphs),
                "table_count": len(tables),
            },
        )