"""
Excel workbook parser for WaterDIP.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from ..base import BaseParser, ParseSource, ParserContext
from ..models import (
    DocumentTable,
    DocumentType,
    WaterDIPDocument,
)
from .common import (
    build_source_provenance,
    deterministic_id,
    filename_title,
)


def serialise_cell_value(
    value: Any,
) -> Any:
    """Convert spreadsheet values into Pydantic/JSON-friendly values."""

    if value is None:
        return None

    if isinstance(
        value,
        (str, int, float, bool),
    ):
        return value

    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except (TypeError, ValueError):
            pass

    return str(value)


class ExcelParser(BaseParser[WaterDIPDocument]):
    """Parse modern Microsoft Excel workbooks."""

    name = "excel"
    parser_version = "1.0"

    supported_extensions = frozenset(
        {
            ".xlsx",
            ".xlsm",
        }
    )

    supported_media_types = frozenset(
        {
            (
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            (
                "application/vnd.ms-excel.sheet."
                "macroenabled.12"
            ),
        }
    )

    priority = 30

    def parse(
        self,
        source: ParseSource,
        context: ParserContext,
    ) -> WaterDIPDocument:
        content = source.read_bytes()

        workbook = load_workbook(
            filename=BytesIO(content),
            read_only=True,
            data_only=True,
        )

        checksum = source.checksum()

        tables: list[DocumentTable] = []
        text_parts: list[str] = []

        try:
            for sheet_index, worksheet in enumerate(
                workbook.worksheets
            ):
                values: list[list[Any]] = []

                for row in worksheet.iter_rows(
                    values_only=True
                ):
                    serialised = [
                        serialise_cell_value(value)
                        for value in row
                    ]

                    if any(
                        value not in (None, "")
                        for value in serialised
                    ):
                        values.append(serialised)

                if not values:
                    continue

                columns = [
                    (
                        str(value).strip()
                        if value not in (None, "")
                        else f"column_{index + 1}"
                    )
                    for index, value in enumerate(values[0])
                ]

                rows = values[1:]

                table = DocumentTable(
                    id=deterministic_id(
                        "wd-table",
                        checksum,
                        worksheet.title,
                    ),
                    title=worksheet.title,
                    columns=columns,
                    rows=rows,
                    metadata={
                        "sheet_index": sheet_index,
                        "sheet_name": worksheet.title,
                        "row_count": len(rows),
                        "column_count": len(columns),
                    },
                )

                tables.append(table)

                text_parts.append(
                    f"Sheet: {worksheet.title}"
                )

                text_parts.append(
                    " | ".join(columns)
                )

                for row in rows:
                    text_parts.append(
                        " | ".join(
                            "" if value is None else str(value)
                            for value in row
                        )
                    )

                text_parts.append("")

        finally:
            workbook.close()

        provenance = build_source_provenance(
            source,
            context,
            parser_name=self.name,
            parser_version=self.parser_version,
            checksum=checksum,
            metadata={
                "sheet_count": len(tables),
            },
        )

        return WaterDIPDocument.create(
            source=provenance,
            title=filename_title(source.filename),
            text="\n".join(text_parts).strip(),
            document_type=DocumentType.SPREADSHEET,
            tables=tables,
            metadata={
                "sheet_count": len(tables),
                "sheet_names": [
                    table.title
                    for table in tables
                ],
            },
        )