from __future__ import annotations

import csv
import json
from io import BytesIO, StringIO
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from app.data_platform.parsing import (
    CSVParser,
    DOCXParser,
    DocumentType,
    ExcelParser,
    JSONParser,
    ParseSource,
    ParserContext,
    ParserRegistry,
    TextParser,
    WaterDIPDocument,
    XMLParser,
    ZIPParser,
    register_builtin_parsers,
)


def parse_bytes(
    parser,
    content: bytes,
    filename: str,
) -> WaterDIPDocument:
    source = ParseSource.from_bytes(
        content,
        filename=filename,
        dataset_id="test-dataset",
    )

    result = parser.execute(
        source,
        ParserContext(),
    )

    return result.document


def test_text_parser() -> None:
    document = parse_bytes(
        TextParser(),
        b"Reservoir pressure analysis.\nWater injection.",
        "reservoir_notes.txt",
    )

    assert document.document_type == DocumentType.TEXT
    assert "Reservoir pressure" in document.text
    assert document.source.parser_name == "text"
    assert len(document.sections) == 1


def test_csv_parser() -> None:
    buffer = StringIO()

    writer = csv.writer(buffer)
    writer.writerow(
        [
            "DATE",
            "OIL_RATE",
            "WATER_RATE",
        ]
    )
    writer.writerow(
        [
            "2026-01-01",
            "1000",
            "500",
        ]
    )

    document = parse_bytes(
        CSVParser(),
        buffer.getvalue().encode("utf-8"),
        "production.csv",
    )

    assert document.document_type == DocumentType.TABULAR_DATA
    assert len(document.tables) == 1

    table = document.tables[0]

    assert table.columns == [
        "DATE",
        "OIL_RATE",
        "WATER_RATE",
    ]

    assert table.rows[0][1] == "1000"


def test_excel_parser() -> None:
    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Production"

    worksheet.append(
        [
            "DATE",
            "OIL_RATE",
            "WATER_RATE",
        ]
    )

    worksheet.append(
        [
            "2026-01-01",
            1000.0,
            500.0,
        ]
    )

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    document = parse_bytes(
        ExcelParser(),
        buffer.getvalue(),
        "production.xlsx",
    )

    assert document.document_type == DocumentType.SPREADSHEET
    assert len(document.tables) == 1

    assert document.tables[0].title == "Production"
    assert document.tables[0].rows[0][1] == 1000.0


def test_docx_parser() -> None:
    source_document = Document()

    source_document.add_heading(
        "Reservoir Engineering",
        level=1,
    )

    source_document.add_paragraph(
        "Reservoir pressure remains stable."
    )

    table = source_document.add_table(
        rows=2,
        cols=2,
    )

    table.cell(0, 0).text = "Parameter"
    table.cell(0, 1).text = "Value"

    table.cell(1, 0).text = "Pressure"
    table.cell(1, 1).text = "3500 psi"

    buffer = BytesIO()
    source_document.save(buffer)

    document = parse_bytes(
        DOCXParser(),
        buffer.getvalue(),
        "reservoir_report.docx",
    )

    assert "Reservoir pressure remains stable." in document.text
    assert len(document.sections) >= 1
    assert len(document.tables) == 1

    assert document.tables[0].rows[0] == [
        "Pressure",
        "3500 psi",
    ]


def test_json_parser() -> None:
    payload = {
        "field": "Volve",
        "well": "15/9-F-1 C",
        "water_cut": 0.35,
    }

    document = parse_bytes(
        JSONParser(),
        json.dumps(payload).encode("utf-8"),
        "engineering.json",
    )

    assert document.document_type == DocumentType.STRUCTURED_DATA

    assert document.metadata["structured_data"]["field"] == "Volve"

    assert '"water_cut": 0.35' in document.text


def test_xml_parser() -> None:
    content = b"""
    <reservoir>
        <field>Volve</field>
        <formation>Hugin</formation>
        <pressure unit="psi">3500</pressure>
    </reservoir>
    """

    document = parse_bytes(
        XMLParser(),
        content,
        "reservoir.xml",
    )

    assert document.document_type == DocumentType.STRUCTURED_DATA

    assert document.metadata["xml_root_tag"] == "reservoir"

    assert "Volve" in document.text
    assert "Hugin" in document.text


def test_zip_parser() -> None:
    buffer = BytesIO()

    with ZipFile(
        buffer,
        mode="w",
        compression=ZIP_DEFLATED,
    ) as archive:
        archive.writestr(
            "reports/report.txt",
            "Reservoir engineering report",
        )

        archive.writestr(
            "production/data.csv",
            "DATE,OIL_RATE\n2026-01-01,1000\n",
        )

    document = parse_bytes(
        ZIPParser(),
        buffer.getvalue(),
        "engineering_package.zip",
    )

    assert document.document_type == DocumentType.ARCHIVE
    assert len(document.attachments) == 2

    filenames = {
        attachment.filename
        for attachment in document.attachments
    }

    assert "reports/report.txt" in filenames
    assert "production/data.csv" in filenames


def test_builtin_parser_registration() -> None:
    registry = ParserRegistry()

    register_builtin_parsers(registry)

    names = set(registry.names())

    assert names == {
        "las",
        "pdf",
        "csv",
        "excel",
        "docx",
        "json",
        "xml",
        "zip",
        "text",
    }


def test_registry_selects_csv() -> None:
    registry = ParserRegistry()

    register_builtin_parsers(registry)

    source = ParseSource.from_bytes(
        b"A,B\n1,2\n",
        filename="data.csv",
    )

    parser = registry.select(source)

    assert parser.name == "csv"